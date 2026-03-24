import os
import re
import time
import threading
import subprocess
import numpy as np
import tkinter as tk
from tkinter import filedialog, messagebox
from PIL import Image
from pdf2image import convert_from_path
import openpyxl
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

# Пути
PHOTOSHOP_EXE = r"D:\Adobe\Adobe Photoshop 2025\Photoshop.exe"
SCRIPT_PATH = r"D:\PyCharm 2024.2.3\PycharmProjects\fillpercent\blackfill.jsx"
OUTPUT_FOLDER = r"D:\PyCharm 2024.2.3\PycharmProjects\fillpercent\Prints"
EXCEL_PATH = "fill_results.xlsx"
POPPER_PATH = r"C:\Program Files\poppler-24.08.0\Library\bin"  # Поправленный путь к Poppler

# Исключаемые слова
EXCLUDE_WORDS = {"processed", "spot", "копия", "новый", "new", "размер", "мм", "сжатие", "актуально", "новая"}

def clean_filename(filename):
    """Очищает имя файла от ненужных слов и цифр."""
    name = os.path.splitext(filename)[0]  # Убираем .pdf
    words = re.findall(r"[а-яА-Яa-zA-Z]+", name)  # Оставляем только слова

    # Убираем ненужные слова + исключаем одиночные "х" и "x"
    filtered_words = [word for word in words if word.lower() not in EXCLUDE_WORDS and word.lower() not in {"х", "x"}]

    return "_".join(filtered_words)  # Объединяем слова через "_"

def extract_images_from_pdf(pdf_path):
    """Извлечение изображений из PDF через pdf2image с отладкой."""
    print(f"🔄 Открываем PDF: {pdf_path}")
    try:
        images = convert_from_path(pdf_path, dpi=300, poppler_path=POPPER_PATH)  # Используем Poppler
        if not images:
            print(f"⚠ Внимание: В PDF {pdf_path} не найдено изображений!")
        return images
    except Exception as e:
        print(f"❌ Ошибка при конвертации PDF {pdf_path}: {e}")
        return []

def calculate_fill_percentage(image: Image):
    """Рассчитать процент заполненности пикселями с отладкой."""
    try:
        img = image.convert("RGB")
        pixels = np.array(img)

        print(f"📊 Обрабатываем изображение: {img.size}, mode: {img.mode}, shape: {pixels.shape}")

        grayscale = np.mean(pixels, axis=2)  # Преобразуем в градации серого
        black_pixels = np.sum(grayscale < 30)  # Чёрный цвет – пиксели с яркостью < 30

        total_pixels = pixels.shape[0] * pixels.shape[1]
        fill_percentage = (black_pixels / total_pixels) * 100
        print(f"✅ Заполненность: {fill_percentage:.2f}%")
        return fill_percentage
    except Exception as e:
        print(f"❌ Ошибка при вычислении заполненности: {e}")
        return 0

def get_image_size_mm(image: Image, dpi=300):
    """Возвращает размеры изображения в миллиметрах."""
    width_mm = (image.width / dpi) * 25.4
    height_mm = (image.height / dpi) * 25.4
    return f"{width_mm:.2f}x{height_mm:.2f} мм"

def save_to_excel(data):
    """Сохранение результатов в Excel."""
    print("💾 Сохраняем данные в Excel...")
    wb = openpyxl.Workbook() if not os.path.exists(EXCEL_PATH) else openpyxl.load_workbook(EXCEL_PATH)
    ws = wb.active if "Results" not in wb.sheetnames else wb["Results"]

    # Заголовки таблицы
    if ws.max_row == 1:
        ws.append(["PDF файл", "Заполненность (%)", "Размер изображения (мм)"])
        for col in range(1, 4):
            ws[f"{get_column_letter(col)}1"].font = Font(bold=True)
            ws[f"{get_column_letter(col)}1"].alignment = Alignment(horizontal="center")

    for row in data:
        ws.append(row)

    wb.save(EXCEL_PATH)
    print(f"✅ Данные сохранены в {EXCEL_PATH}")

def process_pdf(pdf_path):
    """Обрабатывает PDF, извлекает изображения, считает заполненность и сохраняет в Excel."""
    print(f"🔍 Начинаем обработку PDF: {pdf_path}")
    try:
        images = extract_images_from_pdf(pdf_path)
        if not images:
            messagebox.showerror("Ошибка", f"В PDF {os.path.basename(pdf_path)} не найдено изображений!")
            return

        total_fill_percentage = sum(calculate_fill_percentage(img) for img in images)
        avg_fill = total_fill_percentage / len(images) if images else 0
        image_size = get_image_size_mm(images[0]) if images else "Неизвестно"

        clean_name = clean_filename(os.path.basename(pdf_path))  # Очистка названия

        save_to_excel([[clean_name, avg_fill, image_size]])  # Сохраняем в Excel только одно значение

    except Exception as e:
        messagebox.showerror("Ошибка", f"Ошибка при обработке PDF {os.path.basename(pdf_path)}: {e}")

def process_all_pdfs():
    """Выбор папки или файлов для обработки в Photoshop."""
    choice = messagebox.askyesno("Выбор режима", "Обработать всю папку? (Да - выбрать папку, Нет - выбрать файлы)")

    if choice:
        folder_selected = filedialog.askdirectory(title="Выберите папку с PDF-файлами")
        if not folder_selected:
            return

        command = f'start "" "{PHOTOSHOP_EXE}" /r "{SCRIPT_PATH}" "{folder_selected}"'
        print("📤 Передаём команду в Photoshop:", command)
        subprocess.Popen(command, shell=True)

    else:
        file_paths = filedialog.askopenfilenames(filetypes=[("PDF files", "*.pdf")], title="Выберите PDF-файлы")
        if not file_paths:
            return

        def process_files():
            for pdf_path in file_paths:
                pdf_path = os.path.abspath(pdf_path)
                command = f'start "" "{PHOTOSHOP_EXE}" /r "{SCRIPT_PATH}" "{pdf_path}"'
                print("📤 Передаём команду в Photoshop:", command)
                subprocess.Popen(command, shell=True)

                # Ожидание появления обработанного файла
                processed_pdf = os.path.join(OUTPUT_FOLDER, os.path.basename(pdf_path))
                timeout, start_time = 60, time.time()

                while time.time() - start_time < timeout:
                    if os.path.exists(processed_pdf):
                        process_pdf(processed_pdf)
                        break
                    time.sleep(2)

                if not os.path.exists(processed_pdf):
                    messagebox.showerror("Ошибка", f"Файл {os.path.basename(pdf_path)} не появился в папке!")

        threading.Thread(target=process_files, daemon=True).start()

def calculate_percentages():
    """Позволяет выбрать несколько PDF-файлов и рассчитать их заполненность."""
    file_paths = filedialog.askopenfilenames(filetypes=[("PDF files", "*.pdf")], title="Выберите обработанные PDF-файлы")
    if not file_paths:
        return

    def process_selected_files():
        for pdf_path in file_paths:
            process_pdf(pdf_path)

    threading.Thread(target=process_selected_files, daemon=True).start()

def main():
    """Графический интерфейс."""
    root = tk.Tk()
    root.title("PDF Analyzer")
    root.geometry("400x350")
    root.configure(bg="#2E2E2E")

    tk.Button(root, text="Выбрать и обработать PDF", command=process_all_pdfs, bg="#555555", fg="white").pack(pady=10)
    tk.Button(root, text="Посчитать проценты", command=calculate_percentages, bg="#008000", fg="white").pack(pady=10)
    tk.Button(root, text="Выход", command=root.destroy, bg="#8B0000", fg="white").pack(pady=10)

    root.mainloop()

if __name__ == "__main__":
    main()
